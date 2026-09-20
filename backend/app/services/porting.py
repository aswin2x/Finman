"""CSV and Excel import / export for transactions.

Imported rows are marked `is_imported` and tagged with a batch id so a bad
import can be reviewed or rolled back, and every row stays editable.
"""
from __future__ import annotations

import csv
import io
import uuid
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Category, Transaction
from app.services.access import category_visibility_filter
from app.services.finance import money

COLUMNS = ["date", "type", "title", "amount", "category", "payment_method", "scope", "notes"]

DATE_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y", "%d %b %Y", "%d %B %Y")

VALID_TYPES = {"expense", "income"}
VALID_METHODS = {"upi", "cash", "card", "bank", "wallet", "other"}
VALID_SCOPES = {"shared", "personal"}


class ImportError_(Exception):
    pass


def parse_date(raw: object) -> date:
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    text = str(raw).strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ImportError_(f"Unrecognised date: {raw!r}")


def parse_amount(raw: object) -> Decimal:
    text = str(raw).strip().replace(",", "").replace("₹", "").replace("INR", "").strip()
    if text.startswith("(") and text.endswith(")"):
        text = text[1:-1]
    try:
        value = money(Decimal(text))
    except (InvalidOperation, ValueError) as exc:
        raise ImportError_(f"Unrecognised amount: {raw!r}") from exc
    if value <= 0:
        raise ImportError_(f"Amount must be greater than zero: {raw!r}")
    return value


def rows_from_csv(content: bytes) -> list[dict]:
    text = content.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    if reader.fieldnames is None:
        raise ImportError_("The file has no header row")
    return [{(k or "").strip().lower(): v for k, v in row.items()} for row in reader]


def rows_from_xlsx(content: bytes) -> list[dict]:
    wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    try:
        header = [str(h or "").strip().lower() for h in next(rows)]
    except StopIteration as exc:
        raise ImportError_("The sheet is empty") from exc
    out = []
    for raw in rows:
        if raw is None or all(cell is None or str(cell).strip() == "" for cell in raw):
            continue
        out.append({header[i]: raw[i] for i in range(min(len(header), len(raw)))})
    wb.close()
    return out


def preview_rows(rows: list[dict]) -> tuple[list[dict], list[dict]]:
    """Split parsed rows into (valid, errors) without touching the database."""
    valid: list[dict] = []
    errors: list[dict] = []
    for index, row in enumerate(rows, start=2):
        try:
            title = str(row.get("title") or row.get("description") or "").strip()
            if not title:
                raise ImportError_("Missing title")
            txn_type = str(row.get("type") or "expense").strip().lower()
            if txn_type not in VALID_TYPES:
                raise ImportError_(f"Type must be expense or income, got {txn_type!r}")
            method = str(row.get("payment_method") or "upi").strip().lower()
            scope = str(row.get("scope") or "shared").strip().lower()
            valid.append(
                {
                    "row": index,
                    "occurred_on": parse_date(row.get("date")),
                    "type": txn_type,
                    "title": title[:120],
                    "amount": parse_amount(row.get("amount")),
                    "category": str(row.get("category") or "").strip() or None,
                    "payment_method": method if method in VALID_METHODS else "other",
                    "scope": scope if scope in VALID_SCOPES else "shared",
                    "notes": (str(row.get("notes")).strip() if row.get("notes") else None),
                }
            )
        except ImportError_ as exc:
            errors.append({"row": index, "error": str(exc), "data": {k: str(v) for k, v in row.items()}})
    return valid, errors


def resolve_categories(db: Session, wanted: set[tuple[str, str]], user_id: uuid.UUID) -> dict[tuple[str, str], uuid.UUID]:
    """Match category names case-insensitively, creating any that are missing.

    Keyed by (name, kind) so an income row never lands in an expense category.
    Only categories the importer can see are matched; another member's personal
    category is invisible here and a shared one is created instead.
    """
    pairs = {(name.strip().lower(), kind) for name, kind in wanted if name and name.strip()}
    if not pairs:
        return {}

    existing = db.scalars(
        select(Category).where(
            Category.deleted_at.is_(None),
            category_visibility_filter(Category, user_id),
        )
    ).all()
    mapping = {(c.name.strip().lower(), c.kind): c.id for c in existing}

    for name, kind in pairs - set(mapping):
        category = Category(name=name.title(), kind=kind, color="#8A8F98", icon="tag", owner_id=None)
        db.add(category)
        db.flush()
        mapping[(name, kind)] = category.id
    return mapping


def commit_rows(db: Session, rows: list[dict], user_id: uuid.UUID) -> tuple[int, str]:
    batch_id = uuid.uuid4().hex[:16]
    mapping = resolve_categories(
        db, {(r["category"], r["type"]) for r in rows if r["category"]}, user_id
    )
    for row in rows:
        category_id = mapping.get((row["category"].strip().lower(), row["type"])) if row["category"] else None
        db.add(
            Transaction(
                type=row["type"],
                amount=row["amount"],
                title=row["title"],
                notes=row["notes"],
                occurred_on=row["occurred_on"],
                payment_method=row["payment_method"],
                scope=row["scope"],
                category_id=category_id,
                user_id=user_id,
                is_imported=True,
                import_batch_id=batch_id,
            )
        )
    db.commit()
    return len(rows), batch_id


def export_csv(transactions: list[Transaction]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(COLUMNS)
    for t in transactions:
        writer.writerow(
            [
                t.occurred_on.isoformat(),
                t.type,
                t.title,
                f"{money(t.amount)}",
                t.category.name if t.category else "",
                t.payment_method,
                t.scope,
                t.notes or "",
            ]
        )
    return buffer.getvalue().encode("utf-8-sig")


def export_xlsx(transactions: list[Transaction]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws.append([c.title() for c in COLUMNS])
    for t in transactions:
        ws.append(
            [
                t.occurred_on,
                t.type,
                t.title,
                float(money(t.amount)),
                t.category.name if t.category else "",
                t.payment_method,
                t.scope,
                t.notes or "",
            ]
        )
    widths = [12, 10, 32, 12, 18, 16, 10, 40]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + index)].width = width
    ws.freeze_panes = "A2"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def template_csv() -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(COLUMNS)
    writer.writerow(["2026-09-01", "expense", "Groceries", "1500.00", "Groceries", "upi", "shared", "Weekly shop"])
    writer.writerow(["2026-09-01", "income", "Salary", "71000.00", "Salary", "bank", "shared", ""])
    return buffer.getvalue().encode("utf-8-sig")
